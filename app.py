"""
Análise de Preços de Habitação em Portugal — dashboard interativo com dados
reais e oficiais do INE (Instituto Nacional de Estatística).

Fontes (ver scripts/atualizar_dados.py):
- Vendas de alojamentos por concelho/região (€/m²), 2022-2025.
- Rendas (arrendamento) de novos contratos por concelho/região (€/m²/mês), 2022-2025.
- Índice de preços da habitação (vendas), nacional, trimestral, desde 2009.

O mapa usa a fronteira geográfica dos concelhos (dados/concelhos.geojson,
fonte E-REDES/OpenDataSoft — dados abertos, sem necessidade de chave) e um
cruzamento de códigos (dados/concelhos_geojson_crosswalk.csv) para ligar cada
concelho do INE à sua forma no mapa.

Corre localmente com: python app.py  (abre http://localhost:8050)
"""
import io
import os

import dash
import pandas as pd
import plotly.graph_objects as go
from dash import Input, Output, State, ctx, dcc, html
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PASTA_DADOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dados")

# --- Dados (reais, lidos dos CSVs/GeoJSON gerados por scripts/atualizar_dados.py) ---
df_vendas = pd.read_csv(os.path.join(PASTA_DADOS, "precos_regionais.csv"))
df_arrendamento = pd.read_csv(os.path.join(PASTA_DADOS, "rendas_regionais.csv"))
df_indice = pd.read_csv(os.path.join(PASTA_DADOS, "indice_nacional.csv"))
df_crosswalk = pd.read_csv(os.path.join(PASTA_DADOS, "concelhos_geojson_crosswalk.csv"), dtype={"con_code": str})

with open(os.path.join(PASTA_DADOS, "concelhos.geojson"), encoding="utf-8") as f:
    import json

    GEOJSON_CONCELHOS = json.load(f)


# --- Localização e nomes por concelho (pesquisa, clique no mapa, zoom) -----
def _centroide_de_anel(anel: list) -> tuple:
    """
    Centroide de um anel de coordenadas [lon, lat] (fórmula da área de um
    polígono). Se o anel for degenerado (área ~0, ex.: só 2-3 pontos em linha)
    usa a média simples das coordenadas como reserva.
    """
    area2 = 0.0
    cx = 0.0
    cy = 0.0
    for i in range(len(anel) - 1):
        x0, y0 = anel[i][0], anel[i][1]
        x1, y1 = anel[i + 1][0], anel[i + 1][1]
        cruz = x0 * y1 - x1 * y0
        area2 += cruz
        cx += (x0 + x1) * cruz
        cy += (y0 + y1) * cruz
    if abs(area2) < 1e-12:
        lons = [ponto[0] for ponto in anel]
        lats = [ponto[1] for ponto in anel]
        return (sum(lons) / len(lons), sum(lats) / len(lats))
    area = area2 / 2.0
    return (cx / (6.0 * area), cy / (6.0 * area))


def _maior_anel_exterior(geometria: dict):
    """Anel exterior do maior polígono de uma geometria Polygon/MultiPolygon."""
    tipo = geometria.get("type")
    if tipo == "Polygon":
        aneis = [geometria["coordinates"][0]]
    elif tipo == "MultiPolygon":
        aneis = [poligono[0] for poligono in geometria["coordinates"]]
    else:
        aneis = []
    if not aneis:
        return None

    def _area_aprox(anel):
        return abs(sum(anel[i][0] * anel[i + 1][1] - anel[i + 1][0] * anel[i][1] for i in range(len(anel) - 1)))

    return max(aneis, key=_area_aprox)


def _centroide(geometria: dict):
    """Centroide aproximado (lon, lat) de uma geometria GeoJSON Polygon/MultiPolygon."""
    anel = _maior_anel_exterior(geometria)
    if not anel or len(anel) < 3:
        return None
    return _centroide_de_anel(anel)


CENTROIDE_POR_CON_CODE = {}
for _feature in GEOJSON_CONCELHOS["features"]:
    _con_code = _feature.get("properties", {}).get("con_code")
    _centro = _centroide(_feature["geometry"])
    if _con_code and _centro:
        CENTROIDE_POR_CON_CODE[_con_code] = _centro

_concelhos_base = df_vendas[df_vendas["nivel"] == "Concelho"][["geocod", "regiao"]].drop_duplicates("geocod")
NOME_POR_GEOCOD = dict(zip(_concelhos_base["geocod"], _concelhos_base["regiao"], strict=True))
CON_CODE_POR_GEOCOD = dict(zip(df_crosswalk["geocod"], df_crosswalk["con_code"], strict=True))
GEOCOD_POR_CON_CODE = {con_code: geocod for geocod, con_code in CON_CODE_POR_GEOCOD.items()}
CENTROIDE_POR_GEOCOD = {
    geocod: CENTROIDE_POR_CON_CODE[con_code]
    for geocod, con_code in CON_CODE_POR_GEOCOD.items()
    if con_code in CENTROIDE_POR_CON_CODE
}
CONCELHOS_PESQUISA = sorted(
    (
        {"label": nome, "value": geocod}
        for geocod, nome in NOME_POR_GEOCOD.items()
        if geocod in CENTROIDE_POR_GEOCOD
    ),
    key=lambda opcao: opcao["label"],
)

# --- Configuração dos dois tipos de dado (venda / arrendamento) ---
TIPOS = {
    "Venda": {
        "df": df_vendas,
        "coluna": "preco_m2",
        "unidade": "€/m²",
        "unidade_eixo": "€ / m²",
        "formato_curto": "{:.0f}",
        "formato_hover": ".0f",
        "rotulo_exportacao": "Preço (€/m²)",
        "formato_excel": "#,##0",
    },
    "Arrendamento": {
        "df": df_arrendamento,
        "coluna": "renda_m2",
        "unidade": "€/m²/mês",
        "unidade_eixo": "€ / m² / mês",
        "formato_curto": "{:.2f}",
        "formato_hover": ".2f",
        "rotulo_exportacao": "Renda (€/m²/mês)",
        "formato_excel": "#,##0.00",
    },
}
TIPOS_DISPONIVEIS = list(TIPOS.keys())

ANOS_DISPONIVEIS = sorted(df_vendas["ano"].unique().tolist())
ANO_MAIS_RECENTE = ANOS_DISPONIVEIS[-1]
ANO_MAIS_ANTIGO = ANOS_DISPONIVEIS[0]
NIVEIS_DISPONIVEIS = ["NUTS II", "NUTS III", "Concelho"]
QUARTIS_DISPONIVEIS = sorted(df_vendas["quartil"].dropna().unique().tolist())
QUARTIL_MEDIANA = "2.º quartil"

# --- Cores: claro/escuro, seguindo o método "um valor por papel" -----------
# Sequencial (nível de preço, no mapa e nas barras): um único tom (a cor da
# marca), do mais claro ao mais escuro. Categórico (as 3 séries do índice
# nacional): ordem fixa e validada (ΔE >= 8 entre todos os pares, em ambos os
# temas — ver node scripts/validate_palette.js). Divergente (variação
# homóloga, no mapa): dois polos + cinzento neutro no meio.
CORES_CLARO = {
    "fundo": "#f9f9f7",
    "cartao": "#ffffff",
    "borda": "#e4e2dd",
    "texto": "#2b2b28",
    "texto_suave": "#6b6a65",
    "destaque": "#2f6f4f",
    "destaque_suave": "#e5f0ea",
}
CORES_ESCURO = {
    "fundo": "#0d0d0d",
    "cartao": "#1a1a19",
    "borda": "#2c2c2a",
    "texto": "#f2f2f0",
    "texto_suave": "#c3c2b7",
    # o verde da marca (#2f6f4f) só dá 2.9:1 sobre o fundo escuro (#1a1a19);
    # este tom mais claro dá 5.4:1, mantendo a mesma família de cor.
    "destaque": "#45a06e",
    "destaque_suave": "rgba(69, 160, 110, 0.16)",
}

# Categórico (índice nacional: Total / Novos / Existentes) — primeiras 3 cores
# da paleta categórica validada (ordem fixa, nunca trocada entre séries).
CORES_CATEGORICAS = {
    "claro": ["#2a78d6", "#eb6834", "#1baf7a"],
    "escuro": ["#3987e5", "#d95926", "#199e70"],
}

# Divergente (variação homóloga no mapa) — polos + cinzento neutro no meio.
CORES_DIVERGENTE = {
    "claro": {"negativo": "#2a78d6", "neutro": "#f0efec", "positivo": "#e34948"},
    "escuro": {"negativo": "#3987e5", "neutro": "#383835", "positivo": "#e66767"},
}


# Sequencial vivo (nível de preço, no mapa) — rampa "YlOrRd" (ColorBrewer) de
# 7 tons: mais viva do que um único tom da marca, mas continua monótona em
# claridade (amarelo -> laranja -> vermelho escuro), por isso não é uma
# "rainbow" arbitrária — mantém a leitura de "quanto mais escuro/vermelho,
# mais caro". A mesma rampa serve os dois temas: quem muda entre claro/escuro
# é o próprio estilo do mapa base ("carto-positron"/"carto-darkmatter"), e
# estes tons quentes ficam bem visíveis sobre os dois.
RAMPA_NIVEL_VIVIDA = [
    [0.0, "#ffffb2"],
    [0.16, "#fed976"],
    [0.33, "#feb24c"],
    [0.50, "#fd8d3c"],
    [0.66, "#fc4e2a"],
    [0.83, "#e31a1c"],
    [1.0, "#b10026"],
]


def _cores(tema: str) -> dict:
    return CORES_ESCURO if tema == "escuro" else CORES_CLARO


def _kpis_para_ano(tipo: str, ano: int) -> dict:
    """Calcula os números-resumo (estatísticas nacionais) para um ano e tipo (venda/arrendamento)."""
    cfg = TIPOS[tipo]
    df, coluna = cfg["df"], cfg["coluna"]

    linhas_ano = df[
        (df["ano"] == ano) & (df["quartil"] == QUARTIL_MEDIANA) & (df["nivel"] == "Concelho")
    ].dropna(subset=[coluna])

    nacional = df[(df["ano"] == ano) & (df["quartil"] == QUARTIL_MEDIANA) & (df["nivel"] == "Nacional")]
    valor_nacional = nacional[coluna].iloc[0] if not nacional.empty else None

    valor_nacional_anterior = None
    if ano - 1 in ANOS_DISPONIVEIS:
        nacional_ant = df[(df["ano"] == ano - 1) & (df["quartil"] == QUARTIL_MEDIANA) & (df["nivel"] == "Nacional")]
        if not nacional_ant.empty:
            valor_nacional_anterior = nacional_ant[coluna].iloc[0]

    variacao = None
    if valor_nacional and valor_nacional_anterior:
        variacao = (valor_nacional / valor_nacional_anterior - 1) * 100

    mais_caro = linhas_ano.loc[linhas_ano[coluna].idxmax()] if not linhas_ano.empty else None
    mais_barato = linhas_ano.loc[linhas_ano[coluna].idxmin()] if not linhas_ano.empty else None

    return {
        "valor_nacional": valor_nacional,
        "variacao": variacao,
        "mais_caro": mais_caro,
        "mais_barato": mais_barato,
    }


def _cartao_kpi(titulo: str, valor: str, nota: str = "") -> html.Div:
    filhos = [html.Div(titulo, className="kpi-titulo"), html.Div(valor, className="kpi-valor")]
    if nota:
        filhos.append(html.Div(nota, className="kpi-nota"))
    return html.Div(filhos, className="kpi-cartao")


def _linha_evolucao_nacional(tema: str = "claro") -> go.Figure:
    """
    Evolução do índice nacional de vendas desde 2009, com as 3 séries do INE
    (Total, Novos, Existentes) — cor categórica de ordem fixa, legenda,
    hover unificado com linha de referência (spike) e rótulo no fim de cada
    linha (a "válvula de alívio" exigida quando uma das cores, no modo claro,
    fica abaixo do contraste mínimo de 3:1 face à superfície).
    """
    cores = _cores(tema)
    paleta = CORES_CATEGORICAS[tema]
    categorias = [
        ("Total", paleta[0], True),
        ("Novos", paleta[1], False),
        ("Existentes", paleta[2], False),
    ]

    fig = go.Figure()
    for nome, cor, preencher in categorias:
        df = df_indice[df_indice["categoria"] == nome].sort_values("periodo")
        if df.empty:
            continue
        kwargs = dict(
            x=df["periodo"],
            y=df["indice"],
            mode="lines",
            name=nome,
            line=dict(color=cor, width=2.5 if nome == "Total" else 1.75),
            hovertemplate=f"{nome}: " + "%{y:.1f}<extra></extra>",
        )
        if preencher:
            kwargs["fill"] = "tozeroy"
            kwargs["fillcolor"] = cores["destaque_suave"]
        fig.add_trace(go.Scatter(**kwargs))
        ultimo = df.iloc[-1]
        fig.add_annotation(
            x=ultimo["periodo"],
            y=ultimo["indice"],
            text=f" {nome}",
            showarrow=False,
            xanchor="left",
            font=dict(color=cor, size=11),
        )

    fig.update_layout(
        margin=dict(l=10, r=48, t=10, b=10),
        height=340,
        plot_bgcolor=cores["cartao"],
        paper_bgcolor=cores["cartao"],
        font=dict(color=cores["texto"], family="system-ui, sans-serif"),
        hovermode="x unified",
        showlegend=False,  # os rótulos no fim de cada linha substituem a legenda
        xaxis=dict(
            showgrid=False,
            tickangle=-45,
            nticks=16,
            showspikes=True,
            spikemode="across",
            spikesnap="cursor",
            spikecolor=cores["texto_suave"],
            spikethickness=1,
        ),
        yaxis=dict(showgrid=True, gridcolor=cores["borda"], title="Índice (Base 2015 = 100)"),
        hoverlabel=dict(bgcolor=cores["cartao"], font=dict(color=cores["texto"]), bordercolor=cores["borda"]),
        transition=dict(duration=400, easing="cubic-in-out"),
    )
    return fig


def _dados_filtrados(tipo: str, nivel: str, ano: int, quartil: str) -> pd.DataFrame:
    cfg = TIPOS[tipo]
    df = cfg["df"]
    return df[(df["nivel"] == nivel) & (df["ano"] == ano) & (df["quartil"] == quartil)].dropna(
        subset=[cfg["coluna"]]
    )


def _dados_exportacao(tipo: str, nivel: str, ano: int, quartil: str) -> pd.DataFrame:
    """
    Prepara os dados filtrados para exportação (CSV/Excel): colunas com nomes
    legíveis, sem o código interno do INE (geocod), ordenadas por região.
    """
    cfg = TIPOS[tipo]
    df = _dados_filtrados(tipo, nivel, ano, quartil)[["regiao", "nivel", "ano", "quartil", cfg["coluna"]]].copy()
    df = df.rename(
        columns={
            "regiao": "Região",
            "nivel": "Nível geográfico",
            "ano": "Ano",
            "quartil": "Quartil",
            cfg["coluna"]: cfg["rotulo_exportacao"],
        }
    )
    return df.sort_values("Região").reset_index(drop=True)


def _gerar_excel(df: pd.DataFrame, coluna_valor: str, formato_numero: str) -> bytes:
    """
    Gera um .xlsx "pronto a apresentar" a partir de um DataFrame já preparado
    para exportação: cabeçalho a negrito com fundo verde, colunas com largura
    ajustada ao conteúdo, 1ª linha fixa ao scroll (freeze panes), filtros
    automáticos no cabeçalho e a coluna de valores com separador de milhares.
    """
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Dados")
        ws = writer.sheets["Dados"]

        fundo_cabecalho = PatternFill(start_color="2F6F4F", end_color="2F6F4F", fill_type="solid")
        fonte_cabecalho = Font(bold=True, color="FFFFFF")
        for celula in ws[1]:
            celula.fill = fundo_cabecalho
            celula.font = fonte_cabecalho
            celula.alignment = Alignment(horizontal="center", vertical="center")

        for indice, coluna in enumerate(df.columns, start=1):
            letra = get_column_letter(indice)
            largura = max(len(str(coluna)), df[coluna].astype(str).map(len).max()) + 4
            ws.column_dimensions[letra].width = largura

        indice_valor = df.columns.get_loc(coluna_valor) + 1
        for linha in range(2, ws.max_row + 1):
            ws.cell(row=linha, column=indice_valor).number_format = formato_numero

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    return buffer.getvalue()


def _barras_comparacao(
    tipo: str, nivel: str, ano: int, quartil: str, tema: str = "claro", foco_geocod=None
) -> go.Figure:
    cfg = TIPOS[tipo]
    coluna = cfg["coluna"]
    cores = _cores(tema)
    df = _dados_filtrados(tipo, nivel, ano, quartil).sort_values(coluna, ascending=True)

    # Concelho tem 308 barras — mostra só as 25 mais caras para o gráfico ficar legível
    if nivel == "Concelho" and len(df) > 25:
        df = df.sort_values(coluna, ascending=False).head(25).sort_values(coluna, ascending=True)

    # Cor sólida da marca, como sempre — só ganha um contorno quando o
    # concelho em foco (pesquisado/clicado no mapa) está entre as barras
    # visíveis, para o destacar sem mudar a codificação de cor do gráfico.
    marker = dict(color=cores["destaque"])
    if foco_geocod is not None and "geocod" in df.columns and (df["geocod"] == foco_geocod).any():
        marker["line"] = dict(
            color=[cores["texto"] if g == foco_geocod else "rgba(0,0,0,0)" for g in df["geocod"]],
            width=[2.5 if g == foco_geocod else 0 for g in df["geocod"]],
        )

    fig = go.Figure()
    fig.add_trace(
        go.Bar(
            x=df[coluna],
            y=df["regiao"],
            orientation="h",
            marker=marker,
            hovertemplate="%{y}<br>%{x:" + cfg["formato_hover"] + "} " + cfg["unidade"] + "<extra></extra>",
        )
    )
    altura = max(340, 22 * len(df))
    fig.update_layout(
        margin=dict(l=10, r=10, t=10, b=10),
        height=altura,
        plot_bgcolor=cores["cartao"],
        paper_bgcolor=cores["cartao"],
        font=dict(color=cores["texto"], family="system-ui, sans-serif"),
        xaxis=dict(showgrid=True, gridcolor=cores["borda"], title=cfg["unidade_eixo"]),
        yaxis=dict(showgrid=False),
        hoverlabel=dict(bgcolor=cores["cartao"], font=dict(color=cores["texto"]), bordercolor=cores["borda"]),
        transition=dict(duration=500, easing="cubic-in-out"),
    )
    return fig


def _variacao_por_concelho(tipo: str, ano: int, quartil: str) -> pd.DataFrame:
    """
    Variação homóloga (%) do preço/renda por concelho, face ao ano anterior.
    Devolve geocod/regiao/variacao_pct — vazio (sem levantar erro) quando o
    ano anterior não está disponível (ex.: o primeiro ano da série).
    """
    cfg = TIPOS[tipo]
    coluna = cfg["coluna"]
    if ano - 1 not in ANOS_DISPONIVEIS:
        return pd.DataFrame(columns=["geocod", "regiao", "variacao_pct"])

    atual = _dados_filtrados(tipo, "Concelho", ano, quartil)[["geocod", "regiao", coluna]]
    anterior = _dados_filtrados(tipo, "Concelho", ano - 1, quartil)[["geocod", coluna]].rename(
        columns={coluna: "valor_anterior"}
    )
    fundido = atual.merge(anterior, on="geocod", how="inner")
    fundido = fundido[fundido["valor_anterior"] > 0]
    fundido["variacao_pct"] = (fundido[coluna] / fundido["valor_anterior"] - 1) * 100
    return fundido[["geocod", "regiao", "variacao_pct"]]


def _mapa_concelhos(
    tipo: str, ano: int, quartil: str, tema: str = "claro", modo: str = "nivel", foco_geocod=None
) -> go.Figure:
    cfg = TIPOS[tipo]
    cores = _cores(tema)

    if modo == "variacao":
        df_var = _variacao_por_concelho(tipo, ano, quartil)
        if df_var.empty:
            fig = go.Figure()
            fig.update_layout(
                height=440,
                paper_bgcolor=cores["cartao"],
                plot_bgcolor=cores["cartao"],
                font=dict(color=cores["texto_suave"], family="system-ui, sans-serif"),
                xaxis=dict(visible=False),
                yaxis=dict(visible=False),
                annotations=[
                    dict(
                        text=f"Sem dados de {ano - 1} para calcular a variação homóloga de {ano}.",
                        showarrow=False,
                        font=dict(size=13),
                    )
                ],
            )
            return fig

        df = df_var.merge(df_crosswalk[["geocod", "con_code"]], on="geocod", how="inner")
        limite = max(abs(df["variacao_pct"].min()), abs(df["variacao_pct"].max()), 0.1)
        paleta_div = CORES_DIVERGENTE[tema]
        fig = go.Figure(
            go.Choroplethmap(
                geojson=GEOJSON_CONCELHOS,
                locations=df["con_code"],
                z=df["variacao_pct"],
                featureidkey="properties.con_code",
                zmin=-limite,
                zmax=limite,
                colorscale=[
                    [0, paleta_div["negativo"]],
                    [0.5, paleta_div["neutro"]],
                    [1, paleta_div["positivo"]],
                ],
                marker_line_width=0.3,
                marker_line_color=cores["cartao"],
                colorbar=dict(title="Variação (%)", thickness=14, len=0.8, ticksuffix="%"),
                text=df["regiao"],
                hovertemplate="%{text}<br>%{z:+.1f}% face a " + str(ano - 1) + "<extra></extra>",
            )
        )
    else:
        coluna = cfg["coluna"]
        df = _dados_filtrados(tipo, "Concelho", ano, quartil).merge(
            df_crosswalk[["geocod", "con_code"]], on="geocod", how="inner"
        )
        fig = go.Figure(
            go.Choroplethmap(
                geojson=GEOJSON_CONCELHOS,
                locations=df["con_code"],
                z=df[coluna],
                featureidkey="properties.con_code",
                colorscale=RAMPA_NIVEL_VIVIDA,
                marker_line_width=0.3,
                # Contorno num cinzento neutro (não a cor de fundo): o tom mais
                # claro da rampa viva (#ffffb2) fica quase invisível sobre um
                # fundo quase branco — sem um contorno que se distinga da
                # rampa, os concelhos de valor mais baixo "desapareciam" no
                # tema claro.
                marker_line_color=cores["borda"],
                colorbar=dict(title=cfg["unidade"], thickness=14, len=0.8),
                text=df["regiao"],
                hovertemplate="%{text}<br>%{z:" + cfg["formato_hover"] + "} " + cfg["unidade"] + "<extra></extra>",
            )
        )

    centro = {"lat": 39.6, "lon": -8.2}
    zoom = 5.0
    if foco_geocod is not None and foco_geocod in CENTROIDE_POR_GEOCOD:
        lon_foco, lat_foco = CENTROIDE_POR_GEOCOD[foco_geocod]
        centro = {"lat": lat_foco, "lon": lon_foco}
        zoom = 9.5
        fig.add_trace(
            go.Scattermap(
                lat=[lat_foco],
                lon=[lon_foco],
                mode="markers",
                marker=dict(size=18, color=cores["texto"]),
                hoverinfo="skip",
                showlegend=False,
            )
        )

    fig.update_layout(
        map_style="carto-positron" if tema == "claro" else "carto-darkmatter",
        map_zoom=zoom,
        map_center=centro,
        margin=dict(l=0, r=0, t=0, b=0),
        height=440,
        paper_bgcolor=cores["cartao"],
        hoverlabel=dict(bgcolor=cores["cartao"], font=dict(color=cores["texto"]), bordercolor=cores["borda"]),
        transition=dict(duration=500, easing="cubic-in-out"),
    )
    return fig


app = dash.Dash(__name__, title="Análise de Preços de Habitação em Portugal")
server = app.server  # necessário para o Render (gunicorn aponta para "app:server")

# Aplica o tema guardado (localStorage) ANTES da primeira pintura da página,
# num <script> bloqueante logo no <head> — sem isto, uma visita com o tema
# escuro já guardado mostrava sempre um instante de tema claro antes do
# JavaScript do Dash arrancar e corrigir o atributo "data-theme".
app.index_string = """<!DOCTYPE html>
<html>
    <head>
        <script>
        (function () {
            try {
                var guardado = window.localStorage.getItem("tema-armazenado");
                var tema = guardado ? JSON.parse(guardado) : "claro";
                document.documentElement.setAttribute(
                    "data-theme", tema === "escuro" ? "dark" : "light"
                );
            } catch (erro) {
                // localStorage indisponível (ex.: navegação privada) — fica no tema claro por omissão.
            }
        })();
        </script>
        {%metas%}
        <title>{%title%}</title>
        {%favicon%}
        {%css%}
    </head>
    <body>
        {%app_entry%}
        <footer>
            {%config%}
            {%scripts%}
            {%renderer%}
        </footer>
    </body>
</html>"""

app.layout = html.Div(
    [
        dcc.Store(id="tema-armazenado", storage_type="local", data="claro"),
        dcc.Store(id="concelho-selecionado", data=None),
        html.Div(id="tema-dummy", style={"display": "none"}),
        html.Div(
            [
                html.Div(
                    [
                        html.H1("Análise de Preços de Habitação em Portugal"),
                        html.P(
                            "Dados reais e oficiais do INE (Instituto Nacional de Estatística): "
                            "vendas e arrendamento de alojamentos familiares por concelho, e a "
                            "evolução do índice de preços de venda da habitação desde 2009.",
                            className="subtitulo",
                        ),
                    ]
                ),
                html.Button("🌙 Modo escuro", id="botao-tema", className="botao-tema"),
            ],
            className="cabecalho",
        ),
        html.Div(id="kpis", className="kpis-linha"),
        html.Div(
            [
                html.Div(
                    [
                        html.Label("Tipo"),
                        dcc.Dropdown(
                            id="filtro-tipo",
                            options=[{"label": t, "value": t} for t in TIPOS_DISPONIVEIS],
                            value="Venda",
                            clearable=False,
                        ),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label("Nível geográfico"),
                        dcc.Dropdown(
                            id="filtro-nivel",
                            options=[{"label": n, "value": n} for n in NIVEIS_DISPONIVEIS],
                            value="NUTS II",
                            clearable=False,
                        ),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label("Ano"),
                        dcc.Dropdown(
                            id="filtro-ano",
                            options=[{"label": str(a), "value": a} for a in ANOS_DISPONIVEIS],
                            value=ANO_MAIS_RECENTE,
                            clearable=False,
                        ),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label("Quartil"),
                        dcc.Dropdown(
                            id="filtro-quartil",
                            options=[{"label": q, "value": q} for q in QUARTIS_DISPONIVEIS],
                            value=QUARTIL_MEDIANA,
                            clearable=False,
                        ),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label("Pesquisar concelho"),
                        dcc.Dropdown(
                            id="pesquisa-concelho",
                            options=CONCELHOS_PESQUISA,
                            value=None,
                            placeholder="ex.: Sintra",
                            clearable=True,
                            searchable=True,
                        ),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label(" "),
                        html.Button(
                            "✕ Limpar seleção",
                            id="botao-limpar-selecao",
                            className="botao-tema botao-limpar-selecao",
                            disabled=True,
                        ),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label(" "),
                        html.Button("⬇ Descarregar CSV", id="botao-download", className="botao-download"),
                        dcc.Download(id="download-dados"),
                    ],
                    className="filtro",
                ),
                html.Div(
                    [
                        html.Label(" "),
                        html.Button(
                            "📊 Descarregar Excel",
                            id="botao-download-excel",
                            className="botao-download botao-download-excel",
                        ),
                        dcc.Download(id="download-excel"),
                    ],
                    className="filtro",
                ),
            ],
            className="filtros-linha",
        ),
        html.Div(id="painel-concelho", className="painel-concelho"),
        html.Div(
            [
                html.Div(
                    [
                        html.H2("Mapa por concelho"),
                        dcc.RadioItems(
                            id="filtro-modo-mapa",
                            options=[
                                {"label": " Nível de preço", "value": "nivel"},
                                {
                                    "label": " Variação homóloga (%)",
                                    "value": "variacao",
                                    "disabled": ANO_MAIS_RECENTE - 1 not in ANOS_DISPONIVEIS,
                                },
                            ],
                            value="nivel",
                            className="modo-mapa",
                            inline=True,
                        ),
                    ],
                    className="cabecalho-mapa",
                ),
                dcc.Graph(id="grafico-mapa", config={"displayModeBar": False}),
            ],
            className="cartao-grafico cartao-mapa",
        ),
        html.Div(
            [
                html.Div(
                    [
                        html.H2("Comparação entre regiões"),
                        dcc.Graph(id="grafico-comparacao", config={"displayModeBar": False}),
                    ],
                    className="cartao-grafico",
                ),
                html.Div(
                    [
                        html.H2("Evolução do índice nacional (desde 2009)"),
                        dcc.Graph(id="grafico-indice", config={"displayModeBar": False}),
                    ],
                    className="cartao-grafico",
                ),
            ],
            className="graficos-grelha",
        ),
        html.Footer(
            [
                "Fonte: ",
                html.A("INE — Instituto Nacional de Estatística", href="https://www.ine.pt", target="_blank"),
                ". Projeto de portefólio de Elisama Manuel.",
            ],
            className="rodape",
        ),
    ],
    className="pagina",
)


app.clientside_callback(
    """
    function(tema) {
        document.documentElement.setAttribute('data-theme', tema === 'escuro' ? 'dark' : 'light');
        return '';
    }
    """,
    Output("tema-dummy", "children"),
    Input("tema-armazenado", "data"),
)

app.clientside_callback(
    """
    function(n_clicks, tema_atual) {
        if (!n_clicks) { return window.dash_clientside.no_update; }
        return tema_atual === 'escuro' ? 'claro' : 'escuro';
    }
    """,
    Output("tema-armazenado", "data"),
    Input("botao-tema", "n_clicks"),
    State("tema-armazenado", "data"),
)


@app.callback(Output("botao-tema", "children"), Input("tema-armazenado", "data"))
def _atualizar_texto_botao_tema(tema):
    return "☀️ Modo claro" if tema == "escuro" else "🌙 Modo escuro"


@app.callback(
    Output("filtro-modo-mapa", "options"),
    Output("filtro-modo-mapa", "value"),
    Input("filtro-ano", "value"),
    State("filtro-modo-mapa", "value"),
)
def _atualizar_opcoes_modo_mapa(ano, modo_atual):
    tem_ano_anterior = (ano - 1) in ANOS_DISPONIVEIS
    opcoes = [
        {"label": " Nível de preço", "value": "nivel"},
        {"label": " Variação homóloga (%)", "value": "variacao", "disabled": not tem_ano_anterior},
    ]
    novo_modo = modo_atual if (modo_atual == "nivel" or tem_ano_anterior) else "nivel"
    return opcoes, novo_modo


@app.callback(Output("kpis", "children"), Input("filtro-tipo", "value"), Input("filtro-ano", "value"))
def _atualizar_kpis(tipo, ano):
    cfg = TIPOS[tipo]
    k = _kpis_para_ano(tipo, ano)
    cartoes = []

    valor_nacional = cfg["formato_curto"].format(k["valor_nacional"]) + f" {cfg['unidade']}" if k[
        "valor_nacional"
    ] else "—"
    titulo_nacional = "Renda mediana nacional" if tipo == "Arrendamento" else "Preço mediano nacional"
    nota_variacao = ""
    if k["variacao"] is not None:
        sinal = "+" if k["variacao"] >= 0 else ""
        nota_variacao = f"{sinal}{k['variacao']:.1f}% desde {ano - 1}"
    cartoes.append(_cartao_kpi(titulo_nacional, valor_nacional, nota_variacao))

    rotulo_caro = "Concelho mais caro" if tipo == "Venda" else "Concelho com renda mais alta"
    rotulo_barato = "Concelho mais acessível" if tipo == "Venda" else "Concelho com renda mais baixa"

    if k["mais_caro"] is not None:
        cartoes.append(
            _cartao_kpi(
                rotulo_caro,
                k["mais_caro"]["regiao"],
                cfg["formato_curto"].format(k["mais_caro"][cfg["coluna"]]) + f" {cfg['unidade']}",
            )
        )
    if k["mais_barato"] is not None:
        cartoes.append(
            _cartao_kpi(
                rotulo_barato,
                k["mais_barato"]["regiao"],
                cfg["formato_curto"].format(k["mais_barato"][cfg["coluna"]]) + f" {cfg['unidade']}",
            )
        )
    cartoes.append(_cartao_kpi("Concelhos analisados", "308", "todo o país"))
    return cartoes


@app.callback(
    Output("grafico-comparacao", "figure"),
    Input("filtro-tipo", "value"),
    Input("filtro-nivel", "value"),
    Input("filtro-ano", "value"),
    Input("filtro-quartil", "value"),
    Input("tema-armazenado", "data"),
    Input("concelho-selecionado", "data"),
)
def _atualizar_grafico_comparacao(tipo, nivel, ano, quartil, tema, foco_geocod):
    return _barras_comparacao(tipo, nivel, ano, quartil, tema, foco_geocod)


@app.callback(
    Output("grafico-mapa", "figure"),
    Input("filtro-tipo", "value"),
    Input("filtro-ano", "value"),
    Input("filtro-quartil", "value"),
    Input("filtro-modo-mapa", "value"),
    Input("tema-armazenado", "data"),
    Input("concelho-selecionado", "data"),
)
def _atualizar_mapa(tipo, ano, quartil, modo, tema, foco_geocod):
    return _mapa_concelhos(tipo, ano, quartil, tema, modo, foco_geocod)


@app.callback(
    Output("concelho-selecionado", "data"),
    Output("pesquisa-concelho", "value"),
    Input("pesquisa-concelho", "value"),
    Input("grafico-mapa", "clickData"),
    Input("botao-limpar-selecao", "n_clicks"),
    prevent_initial_call=True,
)
def _atualizar_concelho_selecionado(geocod_pesquisa, click_data, _n_clicks_limpar):
    """
    Três origens podem mudar o concelho em foco — a caixa de pesquisa, um
    clique no mapa ou o botão "Limpar seleção" — mas só há um sítio (a Store
    "concelho-selecionado") a guardar isso, por isso é um único callback com
    "ctx.triggered_id" a decidir qual delas disparou, em vez de três
    callbacks a competir pelo mesmo Output (o Dash não permite isso).
    """
    origem = ctx.triggered_id
    if origem == "botao-limpar-selecao":
        return None, None
    if origem == "grafico-mapa":
        if not click_data:
            return dash.no_update, dash.no_update
        con_code = click_data["points"][0].get("location")
        geocod = GEOCOD_POR_CON_CODE.get(con_code)
        if geocod is None:
            return dash.no_update, dash.no_update
        return geocod, geocod
    # origem == "pesquisa-concelho"
    return geocod_pesquisa, geocod_pesquisa


@app.callback(Output("botao-limpar-selecao", "disabled"), Input("concelho-selecionado", "data"))
def _atualizar_estado_botao_limpar(geocod):
    return geocod is None


@app.callback(
    Output("painel-concelho", "children"),
    Input("concelho-selecionado", "data"),
    Input("filtro-tipo", "value"),
    Input("filtro-ano", "value"),
    Input("filtro-quartil", "value"),
)
def _atualizar_painel_concelho(geocod, tipo, ano, quartil):
    if not geocod:
        return []

    nome = NOME_POR_GEOCOD.get(geocod, str(geocod))
    cfg = TIPOS[tipo]
    df = cfg["df"]
    linha = df[(df["geocod"] == geocod) & (df["ano"] == ano) & (df["quartil"] == quartil) & (df["nivel"] == "Concelho")]
    valor_texto = (
        cfg["formato_curto"].format(linha[cfg["coluna"]].iloc[0]) + f" {cfg['unidade']}"
        if not linha.empty
        else "Sem dados para este filtro"
    )

    nota_variacao = ""
    df_var = _variacao_por_concelho(tipo, ano, quartil)
    if not df_var.empty:
        linha_var = df_var[df_var["geocod"] == geocod]
        if not linha_var.empty:
            variacao = linha_var["variacao_pct"].iloc[0]
            sinal = "+" if variacao >= 0 else ""
            nota_variacao = f"{sinal}{variacao:.1f}% desde {ano - 1}"

    return [_cartao_kpi(f"📍 {nome}", valor_texto, nota_variacao)]


@app.callback(Output("grafico-indice", "figure"), Input("tema-armazenado", "data"))
def _atualizar_grafico_indice(tema):
    return _linha_evolucao_nacional(tema)


@app.callback(
    Output("download-dados", "data"),
    Input("botao-download", "n_clicks"),
    State("filtro-tipo", "value"),
    State("filtro-nivel", "value"),
    State("filtro-ano", "value"),
    State("filtro-quartil", "value"),
    prevent_initial_call=True,
)
def _descarregar_csv(n_clicks, tipo, nivel, ano, quartil):
    df = _dados_exportacao(tipo, nivel, ano, quartil)
    nome_ficheiro = f"{tipo.lower()}_{nivel.lower().replace(' ', '-')}_{ano}.csv"
    # Excel em português usa a vírgula como separador decimal, por isso espera
    # o ";" como separador de colunas (senão interpreta o ficheiro inteiro
    # como uma única coluna). Aqui alinhamos com esse formato: ";" a separar
    # colunas, "," como separador decimal, e um BOM UTF-8 (utf-8-sig) para
    # que os acentos dos nomes das regiões apareçam corretamente no Excel.
    return dcc.send_data_frame(
        df.to_csv,
        nome_ficheiro,
        index=False,
        sep=";",
        decimal=",",
        encoding="utf-8-sig",
    )


@app.callback(
    Output("download-excel", "data"),
    Input("botao-download-excel", "n_clicks"),
    State("filtro-tipo", "value"),
    State("filtro-nivel", "value"),
    State("filtro-ano", "value"),
    State("filtro-quartil", "value"),
    prevent_initial_call=True,
)
def _descarregar_excel(n_clicks, tipo, nivel, ano, quartil):
    cfg = TIPOS[tipo]
    df = _dados_exportacao(tipo, nivel, ano, quartil)
    conteudo = _gerar_excel(df, cfg["rotulo_exportacao"], cfg["formato_excel"])
    nome_ficheiro = f"{tipo.lower()}_{nivel.lower().replace(' ', '-')}_{ano}.xlsx"
    return dcc.send_bytes(lambda buffer: buffer.write(conteudo), nome_ficheiro)


if __name__ == "__main__":
    app.run(debug=True, port=8050)